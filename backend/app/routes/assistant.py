import base64
import io
import os
from datetime import datetime, timezone

import requests
from docx import Document
from fastapi import APIRouter, Depends, HTTPException
from starlette.concurrency import run_in_threadpool

from app.database import orders_collection, products_collection
from app.models.assistant import ChatRequest, ChatResponse, ReportFile
from app.routes.deals import apply_discount, get_deal_percents
from app.utils.security import get_optional_admin

router = APIRouter(prefix="/assistant", tags=["assistant"])

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
# gemini-3.1-flash-lite: current stable, free-tier-friendly Gemini model
# (the 2.5 line it replaces is being retired). Override via env if needed.
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.1-flash-lite")
GEMINI_URL = f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent"

MAX_CATALOG_ITEMS = 150  # keep the prompt small even if the catalog grows
MAX_HISTORY_TURNS = 8    # only the recent back-and-forth matters for context

# Used to build a real, clickable product page link for each catalog item.
# Falls back to the known production frontend if FRONTEND_URL isn't set.
STORE_BASE_URL = os.getenv("FRONTEND_URL", "https://vault-weaavver.vercel.app").rstrip("/")

# Only ever exposed to the model when the caller is a verified admin (see
# get_optional_admin below) — a regular customer's request never includes this,
# so Gemini has no way to trigger it for them even if they ask.
SALES_REPORT_TOOL = {
    "name": "generate_sales_report",
    "description": (
        "Generates a downloadable Word (.docx) sales report for one calendar month: "
        "total revenue, number of orders, average order value, and revenue broken down "
        "by product. Only call this when the admin explicitly asks for a sales report, "
        "monthly totals, or a revenue summary."
    ),
    "parameters": {
        "type": "object",
        "properties": {
            "month": {
                "type": "integer",
                "description": "1-12. Defaults to the current month if omitted.",
            },
            "year": {
                "type": "integer",
                "description": "e.g. 2026. Defaults to the current year if omitted.",
            },
        },
    },
}


async def _build_catalog_context() -> str:
    """A compact, LLM-readable snapshot of what's actually for sale, so the
    assistant can only ever recommend real products at their real prices —
    never hallucinated titles or prices."""
    lines = []
    deal_percents = await get_deal_percents()
    cursor = products_collection.find(
        {}, {"title": 1, "price": 1, "type": 1, "platform": 1, "description": 1}
    ).limit(MAX_CATALOG_ITEMS)
    async for p in cursor:
        desc = (p.get("description") or "").strip().replace("\n", " ")
        if len(desc) > 140:
            desc = desc[:140].rsplit(" ", 1)[0] + "..."
        link = f"{STORE_BASE_URL}/products/{p['_id']}"
        price = p.get("price", 0)
        percent = deal_percents.get(str(p["_id"]))
        if percent:
            price_text = f"${apply_discount(price, percent):.2f} (ON SALE, {percent}% off, was ${price:.2f})"
        else:
            price_text = f"${price:.2f}"
        lines.append(
            f"- {p.get('type', 'item')} | \"{p.get('title', 'Untitled')}\" "
            f"| {price_text} | {p.get('platform', 'N/A')} | {desc} | link: {link}"
        )
    return "\n".join(lines) if lines else "(the catalog is currently empty)"


def _system_prompt(catalog: str, is_admin: bool) -> str:
    prompt = (
        "You are the shopping assistant for Vault, a digital game store. "
        "You help visitors find games to buy.\n\n"
        "Rules:\n"
        "- Only recommend or mention products that appear in the CATALOG below. Never "
        "invent titles, prices, or platforms that aren't listed there.\n"
        "- Catalog prices are current and authoritative — never guess or estimate a price. "
        "Entries marked ON SALE are discounted right now: quote the sale price and "
        "mention the original price and the percentage off.\n"
        "- Each catalog entry includes a real product page link. When a visitor asks for a "
        "link, wants to buy something, or would clearly benefit from one, share that exact "
        "link verbatim (don't shorten, alter, or invent one).\n"
        "- If nothing in the catalog matches what the visitor wants, say so honestly and "
        "either suggest the closest alternatives from the catalog or say there isn't a "
        "good match right now.\n"
        "- Keep answers short and conversational — a couple of sentences, or a short "
        "bulleted list of a few options. Don't dump the whole catalog.\n"
        "- If asked something unrelated to shopping, you can answer briefly, then gently "
        "steer the conversation back to helping them find something to buy.\n\n"
        f"CATALOG:\n{catalog}"
    )
    if is_admin:
        prompt += (
            "\n\nYou're talking to a logged-in store admin, not a customer. In addition to "
            "everything above, you can generate a downloadable monthly sales report (total "
            "revenue, order count, and a per-product breakdown) using the "
            "generate_sales_report tool whenever they ask for sales totals, revenue, or a "
            "sales report. Default to the current month if they don't specify one. Never "
            "state a sales figure yourself — only the tool's result is authoritative."
        )
    return prompt


def _gemini_request(body: dict) -> dict:
    if not GEMINI_API_KEY:
        raise HTTPException(
            status_code=503,
            detail="Assistant is not configured (missing GEMINI_API_KEY).",
        )
    try:
        resp = requests.post(GEMINI_URL, params={"key": GEMINI_API_KEY}, json=body, timeout=25)
    except requests.RequestException:
        raise HTTPException(status_code=502, detail="Couldn't reach the assistant service.")

    if resp.status_code == 429:
        raise HTTPException(
            status_code=429,
            detail="The assistant is a little busy right now — try again in a moment.",
        )
    if not resp.ok:
        try:
            message = resp.json().get("error", {}).get("message", "")
        except ValueError:
            message = ""
        detail = "The assistant couldn't respond right now."
        if message:
            detail += f" ({message})"
        raise HTTPException(status_code=502, detail=detail)
    return resp.json()


def _extract_text(data: dict) -> str:
    try:
        parts = data["candidates"][0]["content"]["parts"]
        text = "".join(p.get("text", "") for p in parts if "text" in p).strip()
        return text or "Sorry, I didn't catch that — could you rephrase?"
    except (KeyError, IndexError):
        # Most likely blocked by safety filters, or an unusual/empty response.
        return "Sorry, I couldn't come up with an answer to that — could you try asking differently?"


def _extract_function_call(data: dict):
    try:
        for part in data["candidates"][0]["content"]["parts"]:
            if "functionCall" in part:
                return part["functionCall"]
    except (KeyError, IndexError):
        pass
    return None


async def _generate_sales_report(month: int, year: int) -> dict:
    """Pulls real orders from MongoDB and builds a .docx report. All numbers
    come straight from the database — the LLM never computes them, it only
    triggers this and relays the summary it gets back."""
    if not 1 <= month <= 12:
        month = datetime.now(timezone.utc).month

    start = datetime(year, month, 1, tzinfo=timezone.utc)
    end = datetime(year + 1, 1, 1, tzinfo=timezone.utc) if month == 12 else datetime(year, month + 1, 1, tzinfo=timezone.utc)

    orders = [
        o async for o in orders_collection.find({"created_at": {"$gte": start, "$lt": end}})
    ]

    total_revenue = sum(o.get("total", 0) for o in orders)
    order_count = len(orders)
    avg_order = total_revenue / order_count if order_count else 0

    product_sales = {}
    for o in orders:
        for item in o.get("items", []):
            key = item.get("title", "Unknown")
            entry = product_sales.setdefault(key, {"qty": 0, "revenue": 0.0})
            entry["qty"] += 1
            entry["revenue"] += item.get("price", 0)
    top_products = sorted(product_sales.items(), key=lambda kv: kv[1]["revenue"], reverse=True)

    month_label = start.strftime("%B %Y")

    doc = Document()
    doc.add_heading(f"Vault Sales Report — {month_label}", level=1)

    doc.add_heading("Summary", level=2)
    summary_table = doc.add_table(rows=1, cols=2)
    summary_table.style = "Light Grid Accent 1"
    hdr = summary_table.rows[0].cells
    hdr[0].text, hdr[1].text = "Metric", "Value"
    for label, value in [
        ("Total Revenue", f"${total_revenue:,.2f}"),
        ("Number of Orders", str(order_count)),
        ("Average Order Value", f"${avg_order:,.2f}"),
    ]:
        cells = summary_table.add_row().cells
        cells[0].text, cells[1].text = label, value

    doc.add_heading("Sales by Product", level=2)
    if top_products:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Product", "Units Sold", "Revenue"
        for title, data in top_products:
            cells = table.add_row().cells
            cells[0].text = title
            cells[1].text = str(data["qty"])
            cells[2].text = f"${data['revenue']:,.2f}"
    else:
        doc.add_paragraph("No orders were placed during this period.")

    doc.add_paragraph()
    doc.add_paragraph(
        f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    content_base64 = base64.b64encode(buffer.getvalue()).decode("utf-8")

    summary = (
        f"{order_count} order{'s' if order_count != 1 else ''} totaling "
        f"${total_revenue:,.2f} in revenue for {month_label}"
        + (f", averaging ${avg_order:,.2f} per order." if order_count else ".")
    )

    return {
        "filename": f"vault-sales-report-{year}-{month:02d}.docx",
        "content_base64": content_base64,
        "summary": summary,
    }


@router.post("/chat", response_model=ChatResponse)
async def chat(payload: ChatRequest, is_admin: bool = Depends(get_optional_admin)):
    if not payload.message.strip():
        raise HTTPException(status_code=400, detail="Message can't be empty")

    catalog = await _build_catalog_context()
    system_prompt = _system_prompt(catalog, is_admin)

    contents = [
        {
            "role": "user" if turn.role == "user" else "model",
            "parts": [{"text": turn.content}],
        }
        for turn in payload.history[-MAX_HISTORY_TURNS:]
    ]
    contents.append({"role": "user", "parts": [{"text": payload.message}]})

    body = {
        "system_instruction": {"parts": [{"text": system_prompt}]},
        "contents": contents,
        "generationConfig": {"temperature": 0.4, "maxOutputTokens": 1024},
    }
    if is_admin:
        body["tools"] = [{"functionDeclarations": [SALES_REPORT_TOOL]}]

    data = await run_in_threadpool(_gemini_request, body)
    function_call = _extract_function_call(data)

    if is_admin and function_call and function_call.get("name") == "generate_sales_report":
        args = function_call.get("args") or {}
        now = datetime.now(timezone.utc)
        try:
            month = int(args.get("month") or now.month)
            year = int(args.get("year") or now.year)
        except (TypeError, ValueError):
            month, year = now.month, now.year

        report = await _generate_sales_report(month, year)

        # Feed the real result back so Gemini can phrase a normal reply —
        # it never sees or handles the file bytes itself.
        # Echo back the model's turn exactly as Gemini sent it. Gemini 3 attaches a
        # `thoughtSignature` to the functionCall part, and rebuilding the part by hand
        # (functionCall only) drops it, which makes the follow-up request fail with
        # "Function call is missing a thought_signature".
        contents.append(data["candidates"][0]["content"])
        contents.append(
            {
                "role": "user",
                "parts": [
                    {
                        "functionResponse": {
                            "name": "generate_sales_report",
                            "response": {"summary": report["summary"]},
                        }
                    }
                ],
            }
        )
        body["contents"] = contents
        follow_up = await run_in_threadpool(_gemini_request, body)

        return ChatResponse(
            reply=_extract_text(follow_up),
            report=ReportFile(filename=report["filename"], content_base64=report["content_base64"]),
        )

    return ChatResponse(reply=_extract_text(data))